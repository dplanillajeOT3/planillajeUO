# -*- -*- utf-8 -*-
"""
INTERFAZ GRAFICA INTEGRADORA: DESCARGA POR LOTES Y MODO MANUAL
--------------------------------------------------------------
Panel Izquierdo: Procesamiento por Lotes (Excel)
Panel Derecho: Entrada Manual / Interactiva para un Paciente e Inserción en Matriz
"""

import os
import re
import sys
import json
import threading
from datetime import datetime, date

import tkinter as tk
from tkinter import ttk, messagebox

import openpyxl
from openpyxl.formula.translate import Translator

# Importamos las funciones base desde tu motor principal
import descargar_coberturas as motor

# ============================================================
# LÓGICA COMPLEMENTARIA DE MATRIZ Y BANCO DE DATOS (MODO MANUAL)
# ============================================================

def cargar_datos_pacientes():
    if os.path.exists(motor.RUTA_DATOS_PACIENTES):
        try:
            with open(motor.RUTA_DATOS_PACIENTES, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            return {}
    return {}

def guardar_datos_pacientes(datos):
    try:
        with open(motor.RUTA_DATOS_PACIENTES, "w", encoding="utf-8") as f:
            json.dump(datos, f, ensure_ascii=False, indent=2)
    except Exception as e:
        print(f"Error guardando base de pacientes: {e}")

def buscar_siguiente_fila_vacia(ws):
    fila = 2
    while True:
        val = ws.cell(row=fila, column=motor.COL_MATRIZ_CEDULA).value
        if val is None or str(val).strip() == "":
            return fila
        fila += 1

def copiar_formulas_si_es_necesario(ws, fila_destino):
    fila_origen = fila_destino - 1
    if fila_origen < 2:
        return
    for col in motor._COLUMNAS_FORMULA_MATRIZ:
        celda_origen = ws.cell(row=fila_origen, column=col)
        celda_destino = ws.cell(row=fila_destino, column=col)
        if celda_origen.value and str(celda_origen.value).startswith("="):
            formula_orig = str(celda_origen.value)
            try:
                formula_trad = Translator(formula_orig, origin=celda_origen.coordinate).translate_formula(celda_destino.coordinate)
                celda_destino.value = formula_trad
            except Exception:
                celda_destino.value = formula_orig

def insertar_en_matriz_excel(info_fila, f_nacimiento, sexo_letra, observacion=""):
    os.makedirs(motor.CARPETA_RESPALDOS_MATRIZ, exist_ok=True)
    if os.path.exists(motor.ARCHIVO_MATRIZ):
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        backup = os.path.join(motor.CARPETA_RESPALDOS_MATRIZ, f"INSTRUCTIVO_{stamp}.xlsx")
        try:
            import shutil
            shutil.copy2(motor.ARCHIVO_MATRIZ, backup)
        except Exception:
            pass

    wb = openpyxl.load_workbook(motor.ARCHIVO_MATRIZ)
    ws = wb[motor.HOJA_MATRIZ]

    fila = buscar_siguiente_fila_vacia(ws)
    copiar_formulas_si_es_necesario(ws, fila)

    fecha_atencion = info_fila.get("fecha_atencion")
    cedula_paciente = info_fila.get("cedula_paciente")
    nombre_paciente = info_fila.get("nombre_paciente")
    institucion = info_fila.get("institucion", "IESS")
    tipo_seguro = info_fila.get("tipo_seguro_beneficiario", motor.TIPO_SEGURO_GENERAL_TITULAR)
    parentesco = info_fila.get("parentesco", motor.PARENTESCO_TITULAR)
    cedula_afiliado = info_fila.get("cedula_afiliado", cedula_paciente)
    nombre_afiliado = info_fila.get("nombre_afiliado", nombre_paciente)

    ws.cell(row=fila, column=1, value=fecha_atencion.strftime("%Y-%m-%d") if isinstance(fecha_atencion, (date, datetime)) else fecha_atencion)
    ws.cell(row=fila, column=2, value=institucion)
    ws.cell(row=fila, column=5, value="10")
    ws.cell(row=fila, column=6, value=tipo_seguro)
    ws.cell(row=fila, column=8, value=cedula_paciente)
    ws.cell(row=fila, column=9, value=nombre_paciente)
    ws.cell(row=fila, column=10, value=f_nacimiento.strftime("%Y-%m-%d") if isinstance(f_nacimiento, (date, datetime)) else f_nacimiento)
    ws.cell(row=fila, column=11, value="FEMENINO" if sexo_letra == "F" else "MASCULINO")
    ws.cell(row=fila, column=13, value=cedula_afiliado)
    ws.cell(row=fila, column=14, value=nombre_afiliado)
    ws.cell(row=fila, column=17, value=parentesco)
    ws.cell(row=fila, column=19, value=observacion)

    wb.save(motor.ARCHIVO_MATRIZ)
    return fila

# ============================================================
# CLASE PRINCIPAL DE LA INTERFAZ (TKINTER)
# ============================================================

class AppCovers(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Sistema de Gestión y Descarga de Coberturas de Salud")
        self.geometry("1100x650")
        self.resizable(True, True)

        self.registros_lotes = []
        self.filas_error_lotes = []
        self.cancelar_lotes = False

        self._crear_componentes()

    def _crear_componentes(self):
        # Contenedor Principal en 2 Columnas
        contenedor = ttk.PanedWindow(self, orient=tk.HORIZONTAL)
        contenedor.pack(fill=tk.BOTH, expand=True, padx=5, pady=5)

        # ------------------------------------------------------------
        # PANEL IZQUIERDO: PROCESAMIENTO POR LOTES
        # ------------------------------------------------------------
        frame_izq = ttk.LabelFrame(contenedor, text=" MODO POR LOTES (Excel) ", padding=10)
        contenedor.add(frame_izq, weight=1)

        # Controles
        frame_ctrls = ttk.Frame(frame_izq)
        frame_ctrls.pack(fill=tk.X, pady=5)

        ttk.Button(frame_ctrls, text="Cargar Excel / Iniciar", command=self.iniciar_lotes).pack(side=tk.LEFT, padx=5)
        self.btn_reintentar = ttk.Button(frame_ctrls, text="Reintentar Errores", command=self.reintentar_errores_lotes, state=tk.DISABLED)
        self.btn_reintentar.pack(side=tk.LEFT, padx=5)

        # Tabla de Estado por Fila
        columnas = ("fila", "cedula", "nombre", "estado")
        self.tabla_lotes = ttk.Treeview(frame_izq, columns=columnas, show="headings", height=15)
        self.tabla_lotes.heading("fila", text="Fila")
        self.tabla_lotes.heading("cedula", text="Cédula")
        self.tabla_lotes.heading("nombre", text="Paciente")
        self.tabla_lotes.heading("estado", text="Estado")

        self.tabla_lotes.column("fila", width=40, anchor="center")
        self.tabla_lotes.column("cedula", width=100, anchor="center")
        self.tabla_lotes.column("nombre", width=180, anchor="w")
        self.tabla_lotes.column("estado", width=100, anchor="center")

        scroll_y = ttk.Scrollbar(frame_izq, orient=tk.VERTICAL, command=self.tabla_lotes.yview)
        self.tabla_lotes.configure(yscroll=scroll_y.set)
        
        self.tabla_lotes.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scroll_y.pack(side=tk.RIGHT, fill=tk.Y)

        # Barra de Progreso Lotes
        self.progreso_var = tk.DoubleVar()
        self.barra_progreso = ttk.Progressbar(frame_izq, variable=self.progreso_var, maximum=100)
        self.barra_progreso.pack(fill=tk.X, pady=5, side=tk.BOTTOM)

        # ------------------------------------------------------------
        # PANEL DERECHO: MODO MANUAL / INTERACTIVO
        # ------------------------------------------------------------
        frame_der = ttk.LabelFrame(contenedor, text=" MODO MANUAL / INTERACTIVO ", padding=10)
        contenedor.add(frame_der, weight=1)

        # Formulario
        f_form = ttk.Frame(frame_der)
        f_form.pack(fill=tk.X, pady=5)

        ttk.Label(f_form, text="Cédula Paciente (*):").grid(row=0, column=0, sticky="e", pady=3)
        self.ent_cedula = ttk.Entry(f_form)
        self.ent_cedula.grid(row=0, column=1, sticky="w", pady=3, padx=5)
        self.ent_cedula.bind("<FocusOut>", self._al_cambiar_cedula_manual)

        ttk.Label(f_form, text="Fecha Atención (*):").grid(row=1, column=0, sticky="e", pady=3)
        self.ent_f_atencion = ttk.Entry(f_form)
        self.ent_f_atencion.insert(0, date.today().strftime("%d-%m-%Y"))
        self.ent_f_atencion.grid(row=1, column=1, sticky="w", pady=3, padx=5)

        ttk.Label(f_form, text="Fecha Nacimiento (*):").grid(row=2, column=0, sticky="e", pady=3)
        self.ent_f_nacimiento = ttk.Entry(f_form)
        self.ent_f_nacimiento.grid(row=2, column=1, sticky="w", pady=3, padx=5)

        ttk.Label(f_form, text="Sexo (*):").grid(row=3, column=0, sticky="e", pady=3)
        f_sexo = ttk.Frame(f_form)
        f_sexo.grid(row=3, column=1, sticky="w", pady=3, padx=5)
        self.sexo_var = tk.StringVar(value="M")
        ttk.Radiobutton(f_sexo, text="M", value="M", variable=self.sexo_var).pack(side=tk.LEFT)
        ttk.Radiobutton(f_sexo, text="F", value="F", variable=self.sexo_var).pack(side=tk.LEFT, padx=5)

        ttk.Label(f_form, text="Observación:").grid(row=4, column=0, sticky="e", pady=3)
        self.ent_obs = ttk.Entry(f_form, width=30)
        self.ent_obs.grid(row=4, column=1, sticky="w", pady=3, padx=5)

        btn_procesar_manual = ttk.Button(frame_der, text="Descargar e Insertar en Matriz", command=self.ejecutar_manual)
        btn_procesar_manual.pack(fill=tk.X, pady=10)

        # Consola / Logs Modo Manual
        ttk.Label(frame_der, text="Registros y Diagnóstico del Proceso:").pack(anchor="w")
        self.txt_log_manual = tk.Text(frame_der, height=12, state=tk.DISABLED, bg="#1e1e1e", fg="#00ff00")
        self.txt_log_manual.pack(fill=tk.BOTH, expand=True, pady=5)

    def log_manual(self, texto):
        self.txt_log_manual.config(state=tk.NORMAL)
        self.txt_log_manual.insert(tk.END, texto + "\n")
        self.txt_log_manual.see(tk.END)
        self.txt_log_manual.config(state=tk.DISABLED)

    # ------------------------------------------------------------
    # EVENTOS Y EJECUCIÓN: MODO POR LOTES
    # ------------------------------------------------------------
    def iniciar_lotes(self):
        if not os.path.exists(motor.ARCHIVO_EXCEL):
            messagebox.showerror("Error", f"No se encontró el archivo Excel: {motor.ARCHIVO_EXCEL}")
            return

        for item in self.tabla_lotes.get_children():
            self.tabla_lotes.delete(item)

        self.registros_lotes = motor.leer_excel(motor.ARCHIVO_EXCEL)
        for i, reg in enumerate(self.registros_lotes, start=1):
            self.tabla_lotes.insert("", tk.END, iid=str(i), values=(i, reg["cedula"], reg["nombre"], "Pendiente"))

        threading.Thread(target=self._hilo_lotes, args=(list(enumerate(self.registros_lotes, start=1)),), daemon=True).start()

    def _actualizar_fila_tabla(self, idx, estado):
        self.tabla_lotes.item(str(idx), values=(idx, self.registros_lotes[idx-1]["cedula"], self.registros_lotes[idx-1]["nombre"], estado))

    def _actualizar_progreso(self, n, total):
        porcentaje = (n / total) * 100
        self.progreso_var.set(porcentaje)

    def _hilo_lotes(self, items):
        os.makedirs(motor.CARPETA_SALIDA, exist_ok=True)
        carpeta_temp = os.path.join(motor.CARPETA_SALIDA, "_descargas_temp")
        os.makedirs(carpeta_temp, exist_ok=True)
        carpeta_diag = os.path.join(motor.CARPETA_SALIDA, "_diagnostico")

        driver_holder = [motor.crear_driver(carpeta_temp)]
        errores, sin_seguro = [], []

        try:
            for n, (idx, reg) in enumerate(items, start=1):
                self.after(0, self._actualizar_fila_tabla, idx, "Procesando...")
                
                filas_matriz = motor.procesar_registro(
                    reg, idx, len(items), driver_holder, carpeta_temp, carpeta_diag,
                    errores, sin_seguro
                )
                
                estado_final = "Con errores" if any(e[1] == reg["cedula"] for e in errores) else "OK"
                if estado_final == "Con errores":
                    if idx not in self.filas_error_lotes:
                        self.filas_error_lotes.append(idx)

                self.after(0, self._actualizar_fila_tabla, idx, estado_final)
                self.after(0, self._actualizar_progreso, n, len(items))

            motor._guardar_reportes_finales(errores, sin_seguro)
            self.after(0, lambda: messagebox.showinfo("Proceso Finalizado", "El procesamiento por lotes ha concluido."))
            
            if self.filas_error_lotes:
                self.after(0, lambda: self.btn_reintentar.config(state=tk.NORMAL))

        finally:
            try:
                driver_holder[0].quit()
            except Exception:
                pass

    def reintentar_errores_lotes(self):
        if not self.filas_error_lotes:
            return
        items_reintentar = [(idx, self.registros_lotes[idx-1]) for idx in self.filas_error_lotes]
        self.filas_error_lotes.clear()
        self.btn_reintentar.config(state=tk.DISABLED)
        threading.Thread(target=self._hilo_lotes, args=(items_reintentar,), daemon=True).start()

    # ------------------------------------------------------------
    # EVENTOS Y EJECUCIÓN: MODO MANUAL
    # ------------------------------------------------------------
    def _al_cambiar_cedula_manual(self, event):
        cedula = re.sub(r"\D", "", self.ent_cedula.get()).zfill(10)
        if len(cedula) == 10:
            pacientes = cargar_datos_pacientes()
            if cedula in pacientes:
                info = pacientes[cedula]
                self.ent_f_nacimiento.delete(0, tk.END)
                self.ent_f_nacimiento.insert(0, info.get("fecha_nacimiento", ""))
                self.sexo_var.set(info.get("sexo", "M"))

    def ejecutar_manual(self):
        cedula = re.sub(r"\D", "", self.ent_cedula.get()).zfill(10)
        if len(cedula) != 10:
            messagebox.showerror("Error", "Cédula debe contener 10 dígitos.")
            return

        f_atencion = motor._parsear_fecha_flexible(self.ent_f_atencion.get())
        f_nacimiento = motor._parsear_fecha_flexible(self.ent_f_nacimiento.get())

        if not f_atencion or not f_nacimiento:
            messagebox.showerror("Error", "Formatos de fecha inválidos (use DD-MM-YYYY).")
            return

        sexo = self.sexo_var.get()
        obs = self.ent_obs.get().strip()

        # Guardar cache de paciente
        pacientes = cargar_datos_pacientes()
        pacientes[cedula] = {
            "fecha_nacimiento": f_nacimiento.strftime("%d-%m-%Y"),
            "sexo": sexo
        }
        guardar_datos_pacientes(pacientes)

        threading.Thread(
            target=self._hilo_manual,
            args=(cedula, f_atencion, f_nacimiento, sexo, obs),
            daemon=True
        ).start()

    def _hilo_manual(self, cedula, f_atencion, f_nacimiento, sexo, obs):
        self.after(0, self.log_manual, f"\n=== Consultando Paciente C.I: {cedula} ===")
        
        reg = {
            "cedula": cedula,
            "nombre": None,
            "fechas": [f_atencion],
            "cedula_padre": None,
            "cedula_titular": None
        }

        os.makedirs(motor.CARPETA_SALIDA, exist_ok=True)
        carpeta_temp = os.path.join(motor.CARPETA_SALIDA, "_descargas_temp")
        os.makedirs(carpeta_temp, exist_ok=True)
        carpeta_diag = os.path.join(motor.CARPETA_SALIDA, "_diagnostico")

        driver_holder = [motor.crear_driver(carpeta_temp)]
        errores, sin_seguro = [], []

        try:
            filas_matriz = motor.procesar_registro(
                reg, 1, 1, driver_holder, carpeta_temp, carpeta_diag,
                errores, sin_seguro, recolectar_matriz=True
            )

            if sin_seguro:
                self.after(0, self.log_manual, "=> PACIENTE SIN SEGURO. No se añade a la matriz.")
                motor._guardar_reportes_finales(errores, sin_seguro)
                return

            if errores or not filas_matriz:
                self.after(0, self.log_manual, "=> ERROR obteniendo cobertura. Revisa la consola o logs.")
                motor._guardar_reportes_finales(errores, sin_seguro)
                return

            info_fila = filas_matriz[0]
            fila_excel = insertar_en_matriz_excel(info_fila, f_nacimiento, sexo, obs)
            
            self.after(0, self.log_manual, f"=> ÉXITO: PDF(s) descargados e insertados en Matriz (Fila {fila_excel}).")

        except Exception as e:
            self.after(0, self.log_manual, f"=> Error crítico: {e}")
        finally:
            try:
                driver_holder[0].quit()
            except Exception:
                pass


if __name__ == "__main__":
    app = AppCovers()
    app.mainloop()
